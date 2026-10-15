from abc import ABC, abstractmethod
from io import BytesIO

from docx import Document
from pypdf import PdfReader


class ResumeExtractionError(ValueError):
    pass


class ResumeTextExtractor(ABC):
    @abstractmethod
    def extract(self, contents: bytes) -> str:
        raise NotImplementedError


class PDFTextExtractor(ResumeTextExtractor):
    def extract(self, contents: bytes) -> str:
        try:
            reader = PdfReader(BytesIO(contents))
            pages: list[str] = []
            for page in reader.pages:
                page_text = page.extract_text() or ""
                links: list[str] = []
                for annotation in page.get("/Annots", []) or []:
                    item = annotation.get_object()
                    action = item.get("/A")
                    action = action.get_object() if action is not None and hasattr(action, "get_object") else action
                    uri = action.get("/URI") if action else None
                    if uri and str(uri) not in page_text:
                        links.append(str(uri))
                pages.append("\n".join([page_text, *links]).strip())
            text = "\n".join(page for page in pages if page).strip()
        except Exception as exc:
            raise ResumeExtractionError("Could not read this PDF. Please upload a valid text-based PDF.") from exc
        if not text:
            raise ResumeExtractionError("Could not extract text from this PDF. OCR is not supported in Phase 3.")
        return text


class DOCXTextExtractor(ResumeTextExtractor):
    def extract(self, contents: bytes) -> str:
        try:
            document = Document(BytesIO(contents))
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            tables = [" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()) for table in document.tables for row in table.rows]
            text = "\n".join(paragraphs + [row for row in tables if row]).strip()
        except Exception as exc:
            raise ResumeExtractionError("Could not read this DOCX file. Please upload a valid document.") from exc
        if not text:
            raise ResumeExtractionError("This document does not contain extractable text.")
        return text


def extractor_for_extension(extension: str) -> ResumeTextExtractor:
    if extension == ".pdf": return PDFTextExtractor()
    if extension == ".docx": return DOCXTextExtractor()
    raise ResumeExtractionError("Only PDF and DOCX resume files are supported.")
