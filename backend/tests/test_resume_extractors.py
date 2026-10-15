from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.resume.extractors import DOCXTextExtractor, PDFTextExtractor, ResumeExtractionError


def docx_bytes(text: str = "Ada Developer\nada@example.test\nSkills\nPython, FastAPI") -> bytes:
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Docker"
    table.cell(0, 1).text = "PostgreSQL"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_bytes(*, blank: bool = False) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if not blank:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 720 Td (Ada Developer) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_docx_extracts_paragraphs_and_tables() -> None:
    text = DOCXTextExtractor().extract(docx_bytes())
    assert "Ada Developer" in text
    assert "Docker | PostgreSQL" in text


def test_pdf_extracts_text_offline() -> None:
    assert "Ada Developer" in PDFTextExtractor().extract(pdf_bytes())


@pytest.mark.parametrize("extractor", [PDFTextExtractor(), DOCXTextExtractor()])
def test_invalid_documents_have_safe_errors(extractor) -> None:
    with pytest.raises(ResumeExtractionError, match="Could not read"):
        extractor.extract(b"not a document")


def test_blank_pdf_reports_ocr_limitation() -> None:
    with pytest.raises(ResumeExtractionError, match="OCR is not supported"):
        PDFTextExtractor().extract(pdf_bytes(blank=True))


def test_empty_docx_is_rejected() -> None:
    output = BytesIO()
    Document().save(output)
    with pytest.raises(ResumeExtractionError, match="does not contain"):
        DOCXTextExtractor().extract(output.getvalue())
